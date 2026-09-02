from io import BytesIO

from openpyxl import load_workbook

from tests.test_vessels import vessel_payload


def _create_voyage_with_payment(client, income_amount=500000):
    vessel_id = client.post("/vessels", json=vessel_payload()).json()["id"]
    fixture_id = client.post(
        "/fixtures",
        json={
            "fixture_type": "voyage_charter",
            "charterer": "Test Charterer",
            "contract_currency": "RUB",
            "freight_rate": 25.0,
            "freight_rate_basis": "per_tonne",
        },
    ).json()["id"]
    voyage_id = client.post(
        "/voyages",
        json={
            "fixture_id": fixture_id,
            "vessel_id": vessel_id,
            "voyage_number": "V-001",
            "load_port": "Ust-Luga",
            "discharge_port": "Rotterdam",
            "start_date": "2026-06-01",
            "cargo_grade": "gasoil",
            "cargo_quantity_mt": 5000,
            "laden": True,
        },
    ).json()["id"]
    client.post(
        "/payments",
        json={
            "vessel_id": vessel_id,
            "voyage_id": voyage_id,
            "cost_category": "income",
            "cost_type_name": "Freight",
            "original_currency": "RUB",
            "original_amount": income_amount,
            "invoice_date": "2026-06-01",
        },
    )
    return voyage_id


def test_voyage_pnl_can_be_exported_to_excel(client):
    voyage_id = _create_voyage_with_payment(client)

    response = client.get(f"/reports/voyage-pnl/{voyage_id}?format=xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    header = [cell.value for cell in sheet[1]]
    data_row = [cell.value for cell in sheet[2]]
    assert "revenue" in header
    assert data_row[header.index("revenue")] == 500000


def test_fleet_pnl_can_be_exported_to_excel(client):
    _create_voyage_with_payment(client)

    response = client.get(
        "/reports/fleet-pnl?start_date=2026-06-01&end_date=2026-06-30&format=xlsx"
    )

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    header = [cell.value for cell in workbook.active[1]]
    assert "net_result" in header


def test_claims_status_report_can_be_exported_to_excel(client):
    voyage_id = _create_voyage_with_payment(client)
    client.post(
        "/claims",
        json={
            "voyage_id": voyage_id,
            "claim_type": "cargo_quantity",
            "counterparty": "Rotterdam Terminal Ltd",
            "amount_claimed": 15000,
            "currency": "USD",
            "date_raised": "2026-06-15",
        },
    )

    response = client.get("/reports/claims-status?format=xlsx")

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    sheet = workbook.active
    header = [cell.value for cell in sheet[1]]
    data_row = [cell.value for cell in sheet[2]]
    assert data_row[header.index("counterparty")] == "Rotterdam Terminal Ltd"
