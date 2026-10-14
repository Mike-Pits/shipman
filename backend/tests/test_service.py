import json
import sqlite3
from datetime import datetime, timedelta

import pytest
from sqlalchemy import create_engine

from app.database import Base
from app.services import backup
from tests.test_vessels import vessel_payload


def _backups_dir(tmp_path):
    return tmp_path / "backups"


def test_operator_can_make_a_verified_backup_with_a_manifest(client, tmp_path):
    client.post("/vessels", json=vessel_payload())

    response = client.post("/service/backups")

    assert response.status_code == 201
    manifest = response.json()
    assert manifest["kind"] == "manual"
    assert manifest["integrity_ok"] is True
    assert manifest["fk_violations"] == []
    assert manifest["row_counts"]["vessels"] == 1
    assert len(manifest["sha256"]) == 64
    assert (_backups_dir(tmp_path) / manifest["name"]).exists()
    assert (_backups_dir(tmp_path) / (manifest["name"] + ".json")).exists()


def test_backup_list_and_status_show_the_latest_backup(client):
    assert client.get("/service/backups").json() == []
    assert client.get("/service/status").json()["last_backup"] is None

    created = client.post("/service/backups").json()

    listed = client.get("/service/backups").json()
    assert [b["name"] for b in listed] == [created["name"]]
    status = client.get("/service/status").json()
    assert status["last_backup"]["name"] == created["name"]
    assert status["backup_count"] == 1
    assert status["last_backup_age_hours"] < 1


def test_restore_brings_back_deleted_data_and_keeps_a_safety_copy(client):
    vessel_id = client.post("/vessels", json=vessel_payload()).json()["id"]
    name = client.post("/service/backups").json()["name"]
    client.delete(f"/vessels/{vessel_id}")
    assert client.get("/vessels").json() == []

    response = client.post(f"/service/backups/{name}/restore", json={"confirm": "RESTORE"})

    assert response.status_code == 200
    assert [v["name"] for v in client.get("/vessels").json()] == ["SP Baltic Trader"]
    kinds = {b["name"]: b["kind"] for b in client.get("/service/backups").json()}
    assert kinds[response.json()["safety_backup"]] == "pre-restore"


def test_restore_requires_the_confirmation_word(client):
    name = client.post("/service/backups").json()["name"]

    response = client.post(f"/service/backups/{name}/restore", json={"confirm": "yes"})

    assert response.status_code == 422


def test_restore_is_refused_when_the_backup_lacks_columns_the_app_needs(client, tmp_path):
    name = client.post("/service/backups").json()["name"]
    path = _backups_dir(tmp_path) / name
    conn = sqlite3.connect(path)
    conn.execute("ALTER TABLE vessels DROP COLUMN flag")
    conn.commit()
    conn.close()

    listed = {b["name"]: b for b in client.get("/service/backups").json()}
    assert listed[name]["restore_blockers"] == ["vessels: flag"]
    response = client.post(f"/service/backups/{name}/restore", json={"confirm": "RESTORE"})
    assert response.status_code == 409
    assert "lacks columns" in response.json()["detail"]


def test_restore_is_refused_when_the_backup_file_changed_after_it_was_made(client, tmp_path):
    name = client.post("/service/backups").json()["name"]
    manifest_path = _backups_dir(tmp_path) / (name + ".json")
    manifest = json.loads(manifest_path.read_text())
    manifest["sha256"] = "0" * 64
    manifest_path.write_text(json.dumps(manifest))

    response = client.post(f"/service/backups/{name}/restore", json={"confirm": "RESTORE"})

    assert response.status_code == 409
    assert "checksum" in response.json()["detail"]


def test_backup_can_be_downloaded_and_is_a_sqlite_file(client):
    name = client.post("/service/backups").json()["name"]

    response = client.get(f"/service/backups/{name}/download")

    assert response.status_code == 200
    assert response.content.startswith(b"SQLite format 3\x00")


def test_downloaded_backup_can_be_imported_on_another_install_and_restored(client):
    vessel_id = client.post("/vessels", json=vessel_payload()).json()["id"]
    name = client.post("/service/backups").json()["name"]
    data = client.get(f"/service/backups/{name}/download").content
    client.delete(f"/vessels/{vessel_id}")

    imported = client.post("/service/backups/import", content=data)

    assert imported.status_code == 201
    assert imported.json()["kind"] == "imported"
    restored = client.post(
        f"/service/backups/{imported.json()['name']}/restore", json={"confirm": "RESTORE"}
    )
    assert restored.status_code == 200
    assert len(client.get("/vessels").json()) == 1


def test_import_rejects_files_that_are_not_sqlite_databases(client):
    assert client.post("/service/backups/import", content=b"hello").status_code == 422
    garbage = b"SQLite format 3\x00" + b"\x01" * 500
    assert client.post("/service/backups/import", content=garbage).status_code == 422
    assert client.get("/service/backups").json() == []


def test_backup_names_are_validated_and_unknown_backups_return_404(client):
    assert client.get("/service/backups/bad name.db/download").status_code == 422
    assert client.get("/service/backups/shipman-nope.db/download").status_code == 404
    assert client.delete("/service/backups/shipman-nope.db").status_code == 404


def test_operator_can_delete_a_backup(client, tmp_path):
    name = client.post("/service/backups").json()["name"]

    assert client.delete(f"/service/backups/{name}").status_code == 204

    assert client.get("/service/backups").json() == []
    assert not (_backups_dir(tmp_path) / (name + ".json")).exists()


def test_compact_keeps_the_data_and_takes_a_safety_backup_first(client):
    for i in range(3):
        client.post("/vessels", json=vessel_payload(imo_number=f"91234{i}0"))

    response = client.post("/service/compact")

    assert response.status_code == 200
    body = response.json()
    assert body["size_after"] <= body["size_before"]
    assert len(client.get("/vessels").json()) == 3
    kinds = {b["name"]: b["kind"] for b in client.get("/service/backups").json()}
    assert kinds[body["safety_backup"]] == "pre-compact"


def test_health_check_reports_a_clean_database(client):
    client.post("/vessels", json=vessel_payload())

    health = client.get("/service/health").json()

    assert health["integrity_ok"] is True
    assert health["fk_violations"] == []
    assert health["row_counts"]["vessels"] == 1
    assert health["journal_mode"]
    assert health["size_bytes"] > 0


def test_health_check_finds_rows_pointing_at_missing_parents(client, tmp_path):
    conn = sqlite3.connect(tmp_path / "test.db")
    conn.execute(
        "INSERT INTO estimate_cargo_lots (estimate_id, grade, quantity_mt, freight_rate) "
        "VALUES (999, 'x', 1, 1)"
    )
    conn.commit()
    conn.close()

    health = client.get("/service/health").json()

    assert {"table": "estimate_cargo_lots", "parent": "voyage_estimates", "rows": 1} in health["fk_violations"]


def test_backup_folder_can_be_changed_and_must_be_writable(client, tmp_path):
    elsewhere = tmp_path / "share" / "shipman-backups"

    response = client.put(
        "/service/settings",
        json={"backup_dir": str(elsewhere), "auto_backup_enabled": False, "keep_daily": 7, "keep_monthly": 6},
    )

    assert response.status_code == 200
    assert response.json()["backup_dir"] == str(elsewhere)
    assert response.json()["auto_backup_enabled"] is False
    name = client.post("/service/backups").json()["name"]
    assert (elsewhere / name).exists()

    blocker = tmp_path / "a-file"
    blocker.write_text("x")
    bad = client.put(
        "/service/settings",
        json={"backup_dir": str(blocker / "sub"), "auto_backup_enabled": True, "keep_daily": 7, "keep_monthly": 6},
    )
    assert bad.status_code == 422


def test_export_workbook_has_a_sheet_per_table_with_the_data(client):
    from io import BytesIO

    from openpyxl import load_workbook

    client.post("/vessels", json=vessel_payload())

    response = client.get("/service/export")

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert "vessels" in workbook.sheetnames
    rows = list(workbook["vessels"].iter_rows(values_only=True))
    assert rows[0][1] == "name"
    assert rows[1][1] == "SP Baltic Trader"


# ------------------------------------------------------------- retention / scheduler (unit)


@pytest.fixture()
def engine(tmp_path):
    eng = create_engine(f"sqlite:///{tmp_path / 'unit.db'}")
    Base.metadata.create_all(bind=eng)
    return eng


def _age(eng, name, created_at):
    path = backup.backup_file(eng, name)
    manifest_path = path.with_name(path.name + ".json")
    manifest = json.loads(manifest_path.read_text())
    manifest["created_at"] = created_at.isoformat(timespec="seconds")
    manifest_path.write_text(json.dumps(manifest))


def test_retention_keeps_one_automatic_backup_per_day_and_never_prunes_manual_ones(engine):
    today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    same_day = [backup.create_backup(engine, "auto")["name"] for _ in range(3)]
    for i, name in enumerate(same_day, start=1):
        _age(engine, name, today + timedelta(seconds=i))
    older_day = backup.create_backup(engine, "auto")["name"]
    _age(engine, older_day, today - timedelta(days=3))
    ancient_auto = backup.create_backup(engine, "auto")["name"]
    _age(engine, ancient_auto, today - timedelta(days=400))
    ancient_manual = backup.create_backup(engine, "manual")["name"]
    _age(engine, ancient_manual, today - timedelta(days=400))

    removed = backup.prune(engine)

    remaining = {b["name"] for b in backup.list_backups(engine)}
    assert remaining == {same_day[2], older_day, ancient_manual}
    assert set(removed) == {same_day[0], same_day[1], ancient_auto}


def test_automatic_backup_runs_only_when_the_last_one_is_over_a_day_old(engine):
    first = backup.auto_backup_if_due(engine)
    assert first is not None and first["kind"] == "auto"
    assert backup.auto_backup_if_due(engine) is None

    _age(engine, first["name"], datetime.now() - timedelta(hours=25))
    assert backup.auto_backup_if_due(engine) is not None


def test_automatic_backup_can_be_switched_off(engine):
    backup.save_settings(engine, {"auto_backup_enabled": False})

    assert backup.auto_backup_if_due(engine) is None
    assert backup.list_backups(engine) == []
